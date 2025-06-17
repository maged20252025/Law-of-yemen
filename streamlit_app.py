import streamlit as st
import streamlit.components.v1 as components
from docx import Document
from docx.shared import Inches
import re
import uuid
import os
import time
import html
import csv
from io import BytesIO

# ----------------------------------------------------
# إعدادات الصفحة الأساسية
# ----------------------------------------------------
st.set_page_config(
    page_title="القوانين اليمنية بآخر تعديلاتها حتى عام 2025م",
    layout="wide",
    initial_sidebar_state="expanded" # لجعل الشريط الجانبي مفتوحاً افتراضياً
)

# ----------------------------------------------------
# ثوابت ومتغيرات عامة
# ----------------------------------------------------
TRIAL_DURATION = 180  # 3 دقائق بالثواني
TRIAL_USERS_FILE = "trial_users.txt"
DEVICE_ID_FILE = "device_id.txt"
ACTIVATED_FILE = "activated.txt"
ACTIVATION_CODES_FILE = "activation_codes.txt"
LAWS_DIR = "laws" # تعريف مجلد القوانين هنا

# ----------------------------------------------------
# دوال المساعدة
# ----------------------------------------------------
def get_device_id():
    if os.path.exists(DEVICE_ID_FILE):
        with open(DEVICE_ID_FILE, "r") as f:
            return f.read().strip()
    new_id = str(uuid.uuid4())
    with open(DEVICE_ID_FILE, "w") as f:
        f.write(new_id)
    return new_id

def get_trial_start(device_id):
    if not os.path.exists(TRIAL_USERS_FILE):
        return None
    with open(TRIAL_USERS_FILE, "r") as f:
        reader = csv.reader(f)
        for row in reader:
            if row and row[0] == device_id:
                return float(row[1])
    return None

def register_trial(device_id):
    if not os.path.exists(TRIAL_USERS_FILE):
        with open(TRIAL_USERS_FILE, "w", newline='') as f:
            pass
    with open(TRIAL_USERS_FILE, "a", newline='') as f:
        writer = csv.writer(f)
        writer.writerow([device_id, time.time()])

def is_activated():
    return os.path.exists(ACTIVATED_FILE)

def activate_app(code):
    if not os.path.exists(ACTIVATION_CODES_FILE):
        return False
    with open(ACTIVATION_CODES_FILE, "r") as f:
        codes = [line.strip() for line in f.readlines()]
    if code in codes:
        codes.remove(code)
        with open(ACTIVATION_CODES_FILE, "w") as f:
            for c in codes:
                f.write(c + "\n")
        with open(ACTIVATED_FILE, "w") as f:
            f.write("activated")
        return True
    return False

def highlight_keywords(text, keywords):
    for kw in keywords:
        text = re.sub(f"({re.escape(kw)})", r"<mark>\1</mark>", text, flags=re.IGNORECASE)
    return text

def export_results_to_word(results, filename="نتائج_البحث.docx"):
    document = Document()
    document.add_heading('نتائج البحث في القوانين اليمنية', level=1)
    
    if not results:
        document.add_paragraph("لم يتم العثور على نتائج للكلمات المفتاحية المحددة.")
    else:
        for i, r in enumerate(results):
            document.add_heading(f"القانون: {r['law']} - المادة: {r['num']}", level=2)
            document.add_paragraph(r['plain'])
            if i < len(results) - 1:
                document.add_page_break() 

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

# ----------------------------------------------------
# وظيفة التطبيق الرئيسية (بعد التفعيل أو بدء التجربة)
# ----------------------------------------------------
def run_main_app():
    # الأزرار العائمة للتمرير
    components.html("""
    <style>
    .scroll-btn {
        position: fixed;
        left: 10px;
        padding: 12px;
        font-size: 24px;
        border-radius: 50%;
        background-color: #c5e1a5; /* لون أخضر فاتح */
        color: black;
        cursor: pointer;
        z-index: 9999;
        border: none;
        box-shadow: 1px 1px 5px #888;
    }
    #scroll-top-btn { bottom: 80px; }
    #scroll-bottom-btn { bottom: 20px; }
    </style>
    <button class='scroll-btn' id='scroll-top-btn' onclick='window.scrollTo({top: 0, behavior: "smooth"});'>⬆️</button>
    <button class='scroll-btn' id='scroll-bottom-btn' onclick='window.scrollTo({top: document.body.scrollHeight, behavior: "smooth"});'>⬇️</button>
    """, height=1)

    # التحقق من وجود مجلد القوانين
    if not os.path.exists(LAWS_DIR):
        st.error(f"⚠️ مجلد '{LAWS_DIR}/' غير موجود. يرجى التأكد من وجود ملفات القوانين.")
        return

    files = [f for f in os.listdir(LAWS_DIR) if f.endswith(".docx")]
    if not files:
        st.warning(f"📂 لا توجد ملفات قوانين في مجلد '{LAWS_DIR}/'.")
        return

    # تهيئة session_state لنتائج البحث وحالة البحث
    if "results" not in st.session_state:
        st.session_state.results = []
    if "search_done" not in st.session_state:
        st.session_state.search_done = False
    
    # متغيرة لتخزين الكلمات المفتاحية والقانون المختار من الشريط الجانبي
    # سيتم استخدام هذه القيم في دورات Streamlit التالية
    current_selected_file = None
    current_keywords = ""

    # استخدام الشريط الجانبي للإعدادات
    with st.sidebar:
        st.header("⚙️ إعدادات البحث")
        current_selected_file = st.selectbox("اختر قانونًا للبحث:", ["الكل"] + files, key="sidebar_file_select")
        current_keywords = st.text_area("الكلمات المفتاحية (افصل بفاصلة):", key="sidebar_keywords_input",
                                help="أدخل الكلمات التي تريد البحث عنها، وافصل بينها بفاصلة إذا كانت أكثر من كلمة.")

        # زر البحث في الشريط الجانبي
        # هذا الزر سيقوم بتشغيل دالة البحث مباشرة
        if st.button("🔍 بدء البحث", key="sidebar_search_button", use_container_width=True):
            if current_keywords:
                kw_list = [k.strip() for k in current_keywords.split(",") if k.strip()]
                results = []
                search_files = files if current_selected_file == "الكل" else [current_selected_file]

                with st.spinner("جاري البحث في القوانين... قد يستغرق الأمر بعض الوقت."):
                    for file in search_files:
                        try:
                            doc = Document(os.path.join(LAWS_DIR, file))
                        except Exception as e:
                            st.warning(f"⚠️ تعذر قراءة الملف {file}: {e}. يرجى التأكد من أنه ملف DOCX صالح.")
                            continue

                        law_name = file.replace(".docx", "")
                        last_article = "غير معروفة"
                        current_article_paragraphs = []

                        for para in doc.paragraphs:
                            txt = para.text.strip()
                            if not txt:
                                continue
                            match = re.match(r"مادة\s*[\(]?\s*(\d+)[\)]?", txt)
                            if match:
                                if current_article_paragraphs:
                                    full_text = "\n".join(current_article_paragraphs)
                                    if any(kw.lower() in full_text.lower() for kw in kw_list):
                                        highlighted = highlight_keywords(full_text, kw_list)
                                        results.append({
                                            "law": law_name,
                                            "num": last_article,
                                            "text": highlighted,
                                            "plain": full_text
                                        })
                                    current_article_paragraphs = []
                                last_article = match.group(1)
                            current_article_paragraphs.append(txt)

                        if current_article_paragraphs:
                            full_text = "\n".join(current_article_paragraphs)
                            if any(kw.lower() in full_text.lower() for kw in kw_list):
                                highlighted = highlight_keywords(full_text, kw_list)
                                results.append({
                                    "law": law_name,
                                    "num": last_article,
                                    "text": highlighted,
                                    "plain": full_text
                                })

                st.session_state.results = results
                st.session_state.search_done = True
                if not results:
                    st.info("لم يتم العثور على نتائج للكلمات المفتاحية المحددة.")
                # هنا، بعد انتهاء البحث، يمكننا إجبار Streamlit على إعادة التشغيل
                # لتحديث الواجهة الرئيسية لعرض النتائج.
                # هذا هو أحد الأماكن القليلة التي قد تحتاج فيها إلى rerun بعد إجراء معالجة طويلة.
                # ولكن إذا أزالناها، فسيتم تحديث النتائج في التفاعل التالي للمستخدم.
                # نظرًا للمشاكل السابقة، دعنا نتجنبها هنا أيضاً.
            else:
                st.warning("يرجى إدخال كلمات مفتاحية للبحث.")
        
        st.markdown("---")
        st.info("تلميح: استخدم الشريط الجانبي لإدخال كلمات البحث واختيار القانون، ثم انقر على 'بدء البحث'.")


    # الواجهة الرئيسية لعرض النتائج وزر التصدير
    st.markdown("<h2 style='text-align: center; color: #388E3C;'>نتائج البحث في القوانين 📚</h2>", unsafe_allow_html=True)
    st.markdown("---")

    # عرض زر التصدير ونتائج البحث فقط إذا تم البحث بالفعل وهناك نتائج
    if st.session_state.get("search_done", False): # تأكد من أن البحث قد تم
        results = st.session_state.results
        unique_laws = sorted(set(r["law"] for r in results))

        st.metric(label="📊 إجمالي النتائج التي تم العثور عليها", value=f"{len(results)}", delta=f"في {len(unique_laws)} قانون/ملف")

        if results: # يظهر فقط إذا كانت هناك نتائج
            export_data = export_results_to_word(results)
            st.download_button(
                label="⬇️ تصدير النتائج إلى Word",
                data=export_data,
                file_name="نتائج_البحث_القوانين_اليمنية.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key="download_button_word_main",
                use_container_width=False
            )
        else:
            st.warning("لا توجد نتائج لتصديرها.")
        
        st.markdown("---")

        if results: # فقط إذا كانت هناك نتائج لعرضها
            selected_law_filter = st.selectbox("فلترة النتائج حسب القانون:", ["الكل"] + unique_laws, key="results_law_filter")
            filtered = results if selected_law_filter == "الكل" else [r for r in results if r["law"] == selected_law_filter]

            for i, r in enumerate(filtered):
                with st.expander(f"📚 **{r['law']}** - المادة رقم: **{r['num']}**", expanded=False):
                    st.markdown(f'''
                    <div style="background-color:#f1f8e9;padding:15px;margin-bottom:5px;border-radius:10px;
                                border:1px solid #c5e1a5;direction:rtl;text-align:right;">
                        <p style="font-size:17px;line-height:1.8;margin-top:0px;">
                            {r["text"]}
                        </p>
                    </div>
                    ''', unsafe_allow_html=True)
                    st.text_area(f"📋 المادة كاملة (اضغط لتحديدها ونسخها):", value=r["plain"], height=200, key=f"plain_text_{r['law']}_{r['num']}_{i}")
        else:
            st.info("لا توجد نتائج لعرضها حاليًا. يرجى إجراء بحث جديد.")

# ----------------------------------------------------
# الدالة الرئيسية لتشغيل التطبيق (مع شاشة التفعيل/التجربة)
# ----------------------------------------------------
def main():
    st.markdown("<h1 style='text-align: center; color: #4CAF50;'>📖 تطبيق القوانين اليمنية ⚖️</h1>", unsafe_allow_html=True)
    st.markdown("<h3 style='text-align: center; color: #66BB6A;'>استكشف أحدث التعديلات حتى عام 2025م</h3>", unsafe_allow_html=True)
    st.divider()

    if is_activated():
        run_main_app()
        return

    st.info("👋 مرحباً بك! يرجى تفعيل التطبيق أو بدء التجربة المجانية للاستفادة من جميع الميزات.")

    with st.container(border=True):
        st.subheader("🔐 تفعيل التطبيق")
        code = st.text_input("أدخل كود التفعيل هنا:", key="activation_code_input", help="أدخل الكود الذي حصلت عليه لتفعيل النسخة الكاملة.")
        if st.button("✅ تفعيل الآن", key="activate_button", use_container_width=True):
            if code and activate_app(code.strip()):
                st.success("✅ تم التفعيل بنجاح! يرجى إعادة تشغيل التطبيق لتطبيق التغييرات.")
                st.stop()
            else:
                st.error("❌ كود التفعيل غير صحيح أو انتهت صلاحيته.")

    st.markdown("---")

    with st.container(border=True):
        st.subheader("⏱️ النسخة التجريبية المجانية")
        device_id = get_device_id()
        trial_start = get_trial_start(device_id)

        if trial_start is None:
            if st.button("🚀 بدء التجربة المجانية (3 دقائق)", key="start_trial_button", use_container_width=True):
                register_trial(device_id)
                st.success("✅ بدأت النسخة التجريبية الآن. لديك 3 دقائق لاستخدام التطبيق.")
                st.warning("يرجى التفاعل مع الصفحة (مثلاً، النقر بالماوس أو التمرير) لتحديث الواجهة وبدء استخدام التطبيق.")

        if trial_start is not None:
            elapsed_time = time.time() - trial_start
            remaining_time = TRIAL_DURATION - elapsed_time

            if remaining_time > 0:
                minutes = int(remaining_time // 60)
                seconds = int(remaining_time % 60)
                st.info(f"⏳ النسخة التجريبية لا تزال نشطة. الوقت المتبقي: {minutes:02d}:{seconds:02d}")
                run_main_app()
            else:
                st.error("❌ انتهت مدة التجربة المجانية لهذا الجهاز. يرجى تفعيل التطبيق للاستمرار في الاستخدام.")

if __name__ == "__main__":
    main()
